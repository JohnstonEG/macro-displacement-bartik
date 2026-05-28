"""
Generate Word (.docx) submission files for Labour Economics revision stage.
===========================================================================

Produces three editable Word documents in the project root:
    title_page.docx
    abstract.docx
    declaration.docx

The abstract text pulls its numeric values from paper/stats.tex via regex, so
the generated abstract.docx always reflects the current pipeline output. If
stats.tex is missing a macro, the placeholder "[?]" is rendered, matching the
manuscript's LaTeX behavior.

Setup:
    pip install python-docx

Usage (from anywhere in the project):
    python code/make_submission_docs.py

Outputs land in D:\\Data\\MacroDisplacement\\ next to the .tex equivalents.

Author: Ethan Johnston
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from paths import DATA_DIR

# ---------------------------------------------------------------------------
# Pull numeric values from paper/stats.tex so nothing is hand-typed
# ---------------------------------------------------------------------------
STATS_PATH = Path(DATA_DIR) / "paper" / "stats.tex"
_stats_text = STATS_PATH.read_text(encoding="utf-8") if STATS_PATH.exists() else ""

def stat(name: str) -> str:
    """Look up a \\newcommand value in stats.tex; return [?] if missing."""
    m = re.search(r"\\newcommand\{\\" + re.escape(name) + r"\}\{([^}]*)\}", _stats_text)
    return m.group(1) if m else "[?]"

# ---------------------------------------------------------------------------
# Content constants
# ---------------------------------------------------------------------------
TITLE = ("Cyclical Demand Shocks and Aggregate Task Composition: "
         "Evidence from a Bartik Instrument")

AUTHOR = "Ethan Johnston"
ORCID = "0009-0000-7599-5346"
EMAIL = "egjohn3@ilstu.edu"
AFFILIATION_LINES = [
    "Department of Economics",
    "Illinois State University",
    "Campus Box 4200",
    "Normal, IL 61790, United States",
]

KEYWORDS = ("Task composition; Routine employment; Business cycles; "
            "Local projections; Bartik instrument; Labor reallocation")
JEL_CODES = "E24, E32, J23, J24, J62"

CREDIT = ("Ethan Johnston: Conceptualization, Data curation, Formal analysis, "
          "Methodology, Software, Visualization, Writing - original draft, "
          "Writing - review and editing.")

FUNDING = ("This research did not receive any specific grant from funding "
           "agencies in the public, commercial, or not-for-profit sectors.")

COMPETING = ("The author has no competing financial or personal interests that "
             "could have appeared to influence the work reported in this paper.")

ABSTRACT = (
    f"I estimate the effect of cyclical labor demand shocks on the task "
    f"composition of the employed workforce using local projections instrumented "
    f"with a Bartik shift-share predictor. Exploiting variation in pre-period "
    f"industry composition across {stat('statsNStates')} states from "
    f"{stat('statsYearMin')} to {stat('statsYearMax')}, I find that a "
    f"one-percentage-point increase in state unemployment reduces the mean "
    f"routine task content of employed workers by "
    f"{stat('statsIvRoutineHZeroBeta')} at impact (first-stage F = "
    f"{stat('statsFsF')}), with the effect persisting and modestly intensifying "
    f"over a five-year horizon rather than reverting. The aggregate shift is "
    f"concentrated in the routine-manual employment share, which falls sharply "
    f"during downturns and remains depressed; the decline is offset by rising "
    f"non-routine employment, with the non-routine-manual share the most "
    f"persistent absorbing margin. Controlling for cumulative changes in "
    f"industry employment composition attenuates the impact effect by only "
    f"{stat('statsMechAttenuationHzero')}%, indicating that the response operates "
    f"primarily within rather than between industries. Placebo tests reveal no "
    f"significant differential pre-trend in routine task content, the labor-"
    f"force-participation response is too small at all horizons to drive the "
    f"headline pattern, and the result survives flexible technology controls, "
    f"state-specific trends, region-by-year trends, leave-one-industry-out "
    f"variants of the Bartik instrument, a drop-COVID subsample, and weak-"
    f"instrument-robust inference. The findings are consistent with cyclical "
    f"demand shocks contributing to long-run task reallocation rather than "
    f"generating transitory deviations around a structural trend."
)

# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------
def set_default_font(doc, font_name: str = "Times New Roman", size_pt: int = 11):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size_pt)

def add_centered(doc, text: str, *, bold: bool = False, size_pt: int = 11,
                 space_after_pt: int = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after_pt)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size_pt)
    return p

def add_bold_label(doc, label: str, body: str):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(" " + body)
    return p

# ---------------------------------------------------------------------------
# 1. title_page.docx
# ---------------------------------------------------------------------------
def make_title_page(out_path: Path):
    doc = Document()
    set_default_font(doc)

    add_centered(doc, TITLE, bold=True, size_pt=14, space_after_pt=24)
    add_centered(doc, AUTHOR, size_pt=12, space_after_pt=6)

    for line in AFFILIATION_LINES:
        add_centered(doc, line, size_pt=11, space_after_pt=0)

    doc.add_paragraph()
    add_centered(doc, f"Email: {EMAIL}", size_pt=11, space_after_pt=0)
    add_centered(doc, f"ORCID: {ORCID}", size_pt=11, space_after_pt=18)

    add_bold_label(
        doc,
        "Corresponding author.",
        f"{AUTHOR}, Department of Economics, Illinois State University, "
        f"Campus Box 4200, Normal, IL 61790, United States. Email: {EMAIL}.",
    )

    doc.add_paragraph()
    add_bold_label(doc, "CRediT authorship contribution statement.", CREDIT)

    doc.add_paragraph()
    add_bold_label(doc, "Funding.", FUNDING)

    doc.save(out_path)
    print(f"  Saved: {out_path}")

# ---------------------------------------------------------------------------
# 2. abstract.docx  (anonymized — no author info, suitable for double-blind
#                    review system that uses this as a separate item)
# ---------------------------------------------------------------------------
def make_abstract(out_path: Path):
    doc = Document()
    set_default_font(doc)

    add_centered(doc, TITLE, bold=True, size_pt=14, space_after_pt=24)

    p = doc.add_paragraph()
    r = p.add_run("Abstract")
    r.bold = True
    r.font.size = Pt(12)

    doc.add_paragraph(ABSTRACT)
    doc.add_paragraph()

    add_bold_label(doc, "Keywords:", KEYWORDS + ".")
    add_bold_label(doc, "JEL Codes:", JEL_CODES + ".")

    doc.save(out_path)
    print(f"  Saved: {out_path}")

# ---------------------------------------------------------------------------
# 3. declaration.docx
# ---------------------------------------------------------------------------
def make_declaration(out_path: Path):
    doc = Document()
    set_default_font(doc)

    add_centered(doc, "Declaration of Competing Interest",
                 bold=True, size_pt=14, space_after_pt=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.italic = True
    r.font.size = Pt(11)

    add_centered(doc, f"{AUTHOR}, Illinois State University",
                 size_pt=11, space_after_pt=24)

    doc.add_paragraph(COMPETING)
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph("_" * 32)
    doc.add_paragraph(AUTHOR)
    doc.add_paragraph("Department of Economics")
    doc.add_paragraph("Illinois State University")
    doc.add_paragraph(f"Email: {EMAIL}")

    doc.save(out_path)
    print(f"  Saved: {out_path}")

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    out_dir = Path(DATA_DIR)
    print("=" * 70)
    print("SUBMISSION DOCUMENT GENERATOR")
    print("=" * 70)
    print(f"Reading stats from: {STATS_PATH}")
    print(f"Output directory:   {out_dir}")
    print()

    make_title_page(out_dir / "title_page.docx")
    make_abstract(out_dir / "abstract.docx")
    make_declaration(out_dir / "declaration.docx")

    print()
    print("=" * 70)
    print("DONE. Three .docx files ready for upload.")
    print("=" * 70)

if __name__ == "__main__":
    main()
